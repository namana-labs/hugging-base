"""res: ERCOT-wide reserves on 25 Sep 2026 (REAL), from saved raw fetches only. No network."""
import json, csv, glob, sys
from datetime import datetime, timezone, timedelta

LIVE = '/Users/rzalagbada/Desktop/projects/base-power-hackathon/evidence/live-20260925/'
CACHE = '/Users/rzalagbada/Desktop/projects/base-power-hackathon/evidence/scratchpad-20260925/'
REL_LIVE = 'evidence/live-20260925/'
REL_CACHE = 'evidence/scratchpad-20260925/'
CDT = timezone(timedelta(hours=-5))

def kv(ascm):
    out = {}
    for grp in ascm['data'].values():
        for k, v in grp[1:]:
            out[k] = v
    return out

# AS Plan NP4-33-CD (published 2026-09-25 05:00, covers 25 Sep..1 Oct) and DAM MCPC NP4-188-CD (OD 25 Sep)
plan = {}; mcpc = {}
for row in csv.DictReader(open(glob.glob(LIVE + 'res-np4-33-asplan-pub20260925/*.csv')[0])):
    if row['DeliveryDate'] == '09/25/2026':
        plan.setdefault(row['AncillaryType'], {})[int(row['HourEnding'][:2])] = float(row['Quantity'])
for row in csv.DictReader(open(glob.glob(LIVE + 'res-np4-188-dam-mcpc-od20260925/*.csv')[0])):
    if row['DeliveryDate'] == '09/25/2026':
        mcpc.setdefault(row['AncillaryType'], {})[int(row['HourEnding'][:2])] = float(row['MCPC'])
assert all(len(plan[p]) == 24 for p in plan) and all(len(mcpc[p]) == 24 for p in mcpc)

PARTS = {
    'RRS': [('pfrGenEsr', 'Gen + ESR (primary frequency response)', 'rrAwdGen', 'rrcCapPfrGenEsr'),
            ('ufrLoad', 'Load Resources ex-CLR (under-frequency relay)', 'rrAwdNonClr', 'rrcCapLrWoClr'),
            ('pfrClr', 'Controllable Load Resources (PFR)', 'rrAwdClr', 'rrcCapLr'),
            ('ffr', 'Fast Frequency Response (award: all; capability: non-ESR + ESR rows)', 'rrAwdFfr', ('rrcCapFfr', 'rrcCapFfrEsr'))],
    'ECRS': [('gen', 'Generation Resources', 'ecrsAwdGen', 'ecrsCapGen'),
             ('nclr', 'Load Resources ex-CLR', 'ecrsAwdNonClr', 'ecrsCapNclr'),
             ('clr', 'Controllable Load Resources', 'ecrsAwdClr', 'ecrsCapClr'),
             ('qs', 'Quick Start Gen (OFFQS)', 'ecrsAwdQs', 'ecrsCapQs'),
             ('esr', 'Energy Storage Resources', 'ecrsAwdEsr', 'ecrsCapEsr')],
    'NSPIN': [('onGen', 'On-line Gen with energy offers (incl. QSGRs)', 'nsrAwdGenWEo', 'nsrCapOnGenWoEo'),
              ('outputSched', 'On-line Gen with output schedules', 'nsrAwdGenWOs', 'nsrCapOffResWOs'),
              ('load', 'Load Resources', 'nsrAwdLr', 'nsrCapUndeployedLr'),
              ('offGen', 'Off-line Gen ex-QSGR (+ power augmentation)', 'nsrAwdOffGen', 'nsrCapOffGen'),
              ('qs', 'Quick Start Gen (OFFQS); no matching capability row', 'nsrAwdQs', None),
              ('esr', 'Energy Storage Resources', 'nsrAwdAs', 'nsrCapEsr')],
}
NAMES = {'REGUP': 'Reg-Up', 'REGDN': 'Reg-Down', 'RRS': 'Responsive Reserve (RRS, incl. FFR)',
         'ECRS': 'ERCOT Contingency Reserve (ECRS)', 'NSPIN': 'Non-Spinning Reserve (Non-Spin)'}
DURATION_H = {'REGUP': 0.5, 'REGDN': 0.5, 'RRS': 0.5, 'ECRS': 1.0, 'NSPIN': 4.0}  # NPRR1282 + ERCOT Board item 15
RESPONSE = {'REGUP': 'every 4 s (AGC)', 'REGDN': 'every 4 s (AGC)', 'RRS': 'seconds (PFR/UFR); FFR within 15 cycles',
            'ECRS': 'within 10 min', 'NSPIN': 'within 30 min'}
UP_PRODUCTS = ('REGUP', 'RRS', 'ECRS', 'NSPIN')   # the AORDC covers every up-reserve product, i.e. all but Reg-Down (SOM footnote 94)
AORDC_MW = 10000.0
SOM = 'Potomac Economics, 2025 State of the Market Report for ERCOT, Appendix A (evidence/scratchpad-20260925/som2025.txt lines 8116-8125, footnote 94 at 8131, and 8144-8152)'
AORDC_SRC = SOM + '; plans: NP4-33-CD doc 1278738466'
CHIP_NSPIN = 'RT Non-Spin exceeds the plan by design (the NSRS ASDC is extended up to the 10,000 MW AORDC; IMM 2025 SOM App. A).'
GROSS_NOTE = ('prcMW, upAnyAsComboMW and upRegRrsEcrsMW are GROSS capability. They include capacity that already backs AS awards, '
              'so none of them is spare reserve above awards. upAnyAsComboMW = AS-capable Gen (incl. ESRs) and Load Resources, '
              '"considering current output level and HSL/LSL" (capacity-monitor page; Protocols 6.5.7.5(1)(xv)(C)). PRC = the '
              'Protocols 6.5.7.5(1)(p) formula: On-Line headroom (each unit capped at 20% of its HSL) plus Load Resource terms.')
SPARE_CAVEAT = ('Approximation only. The row and the award sum use different resource sets: off-line Non-Spin awards (off-line Gen + OFFQS) '
                'are in the award sum but may not be in the row, which considers current output and HSL (UNVERIFIED). If they are '
                'outside it, this understates spare capability by up to offlineNonSpinAwardMW. Not a transmission-netted figure.')
FOOTNOTES = {
    'transmission': ('Awards shown are after any manual transmission derate: ERCOT may reduce "the amount of Ancillary Service eligible to '
                     'be awarded to a Resource" (Protocols 6.4.9.1.1(6)), awards are based on "Ancillary Service limits" (6.4.9.1.1(1)), '
                     'and the next SCED run awards against the new limit (clearing the volume on other Resources or going short). The monitor '
                     'does not show which Resources were derated (posted 60 days later, 6.4.9.1.1(7)) and cannot show reserve behind a '
                     'constraint that ERCOT did not derate.'),
    'distribution': ('ADERs: "Identified limitations on the distribution system will not explicitly be enforced by ERCOT\'s systems in '
                     'awarding or dispatching the ADER" (ADER Governing Document Phase 3.3).'),
    'nonSpin': CHIP_NSPIN,
    'gross': GROSS_NOTE}


def snapshot(path, relpath, retrieved_utc, label):
    raw = json.load(open(path)); v = kv(raw)
    as_of = datetime.strptime(raw['lastUpdated'], '%Y-%m-%d %H:%M:%S%z')
    he = as_of.hour + 1  # hour ending
    prods = []
    for pid in ['REGUP', 'REGDN', 'RRS', 'ECRS', 'NSPIN']:
        p = {'id': pid, 'name': NAMES[pid], 'durationH': DURATION_H[pid], 'response': RESPONSE[pid],
             'planMW': plan[pid][he], 'damMcpcUsdPerMWh': mcpc[pid][he]}
        if pid in ('REGUP', 'REGDN'):
            s = 'Up' if pid == 'REGUP' else 'Down'
            p.update(awardMW=v[f'reg{s}Awd'], capabilityMW=v[f'reg{s}Cap'], deployedMW=v[f'reg{s}Deployed'],
                     undeployedMW=v[f'reg{s}Undeployed'], esrAwardMW=None, esrCapabilityMW=None, parts=None)
        else:
            parts = []
            for key, lab, ak, ck in PARTS[pid]:
                cap = None if ck is None else (sum(v[c] for c in ck) if isinstance(ck, tuple) else v[ck])
                parts.append({'key': key, 'label': lab, 'awardMW': v[ak], 'capabilityMW': cap})
            p['parts'] = parts
            p['awardMW'] = sum(x['awardMW'] for x in parts)
            p['capabilityMW'] = sum(x['capabilityMW'] for x in parts if x['capabilityMW'] is not None)
            esr = next((x for x in parts if x['key'] == 'esr'), None)
            p['esrAwardMW'] = esr['awardMW'] if esr else None
            p['esrCapabilityMW'] = esr['capabilityMW'] if esr else None
            if pid == 'RRS':
                p['esrNote'] = 'ESR share not separable: ERCOT reports PFR from Gen and ESR as one row; FFR-from-ESR capability row = %d MW' % v['rrcCapFfrEsr']
            if pid == 'NSPIN':
                p['chip'] = CHIP_NSPIN
                p['compositionNote'] = ('RT Non-Spin exceeds the AS Plan by design: RT awards follow ASDCs, not the Plan (Protocols 6.4.9.1.1(1)), '
                                        'and the up-reserve ASDCs must add up to the 10,000 MW AORDC; when the four up-reserve Plans sum to less, '
                                        '"This excess volume is currently assigned to the ASDC for NSRS, which causes the real-time market to procure '
                                        'excess NSRS" (IMM 2025 SOM, App. A). See aordc for the DERIVED ceiling per snapshot. ERCOT defines the on-line '
                                        'Gen row as including QSGR awards and also lists a separate OFFQS Quick Start row with no matching capability '
                                        'row, so award minus capability is not like-for-like. ancillary-services.json nsrs sparkline = on-line Gen + '
                                        'Load + off-line Gen rows (see series).')
        p['awardMinusCapabilityMW'] = p['awardMW'] - p['capabilityMW']
        p['esrShareOfAward'] = None if not p.get('esrAwardMW') else round(p['esrAwardMW'] / p['awardMW'], 3)
        prods.append(p)
    byid = {p['id']: p for p in prods}
    up_award = sum(byid[k]['awardMW'] for k in UP_PRODUCTS)
    system = {'prcMW': v['prc'],
              'upAnyAsComboMW': v['sumCapResRegUpRrsEcrsNsr'],
              'upRegRrsEcrsMW': v['sumCapResRegUpRrsEcrs'],
              'esrHeadroomUpWithOffersMW': v['esrCapWEoIncreaseBp'],
              'genHeadroomUpWithOffersMW': v['capWEoIncreaseBp'],
              'genHeadroomUp5minHdlMW': v['capIncreaseGenBp'],
              'outOfServiceHslMW': v['telemHslOut'], 'emrHslMW': v['telemHslEmr'],
              'grossNote': GROSS_NOTE,
              'upAwardedMW': up_award,
              'spareApprox': {'status': 'DERIVED', 'MW': v['sumCapResRegUpRrsEcrsNsr'] - up_award,
                              'formula': f"upAnyAsComboMW - (Reg-Up + RRS + ECRS + Non-Spin RT awards) = {v['sumCapResRegUpRrsEcrsNsr']} - {up_award}",
                              'offlineNonSpinAwardMW': sum(x['awardMW'] for x in byid['NSPIN']['parts'] if x['key'] in ('offGen', 'qs')),
                              'caveat': SPARE_CAVEAT}}
    # AORDC ceiling on RT Non-Spin (DERIVED from the IMM's description of the ASDC extension + the NP4-33-CD Plan)
    plan_up = sum(byid[k]['planMW'] for k in UP_PRODUCTS)
    ext = max(0.0, AORDC_MW - plan_up)
    ceiling = byid['NSPIN']['planMW'] + ext
    over = byid['NSPIN']['awardMW'] - ceiling
    qs = next(x['awardMW'] for x in byid['NSPIN']['parts'] if x['key'] == 'qs')
    aordc = {'status': 'DERIVED', 'aordcMW': AORDC_MW, 'planSumUpMW': plan_up, 'nsrsExtensionMW': ext,
             'nonSpinCeilingMW': ceiling, 'nonSpinAwardMW': byid['NSPIN']['awardMW'],
             'nonSpinOverCeilingMW': max(0.0, over), 'offqsAwardMW': qs,
             'nonSpinOverCeilingIfOffqsOverlapsMW': max(0.0, over - qs),
             'upAwardSumMW': up_award,
             'formula': (f"ceiling = NSPIN plan + max(0, 10,000 - (REGUP + RRS + ECRS + NSPIN plans)) = {byid['NSPIN']['planMW']:.0f} + "
                         f"(10,000 - {plan_up:.0f}) = {ceiling:.0f} MW; plans from NP4-33-CD HE{he}"),
             'source': AORDC_SRC}
    if over > 0:
        aordc['residualNote'] = (f"UNVERIFIED: the award is {over:.0f} MW above this ceiling ({max(0.0, over - qs):.0f} MW if the {qs} MW OFFQS row "
                                 f"overlaps the on-line Gen row that 'includes' QSGRs); all up-reserve awards sum to {up_award} MW, above the "
                                 "10,000 MW AORDC itself. Not established whether ERCOT changed the hour's Plan intraday (Protocols 6.4.9.1.2), "
                                 "or the posted ASDCs differ from this reading; ERCOT's posted ASDCs were not fetched.")
    return {'id': label[0], 'label': label[1], 'asOf': as_of.isoformat(), 'retrievedUtc': retrieved_utc,
            'planHourEnding': he, 'source': relpath, 'status': 'REAL', 'products': prods, 'system': system, 'aordc': aordc}


snaps = [
    snapshot(CACHE + 'ercot/db_ascm.json', REL_CACHE + 'ercot/db_ascm.json', '2026-09-26T00:30:08Z',
             ('peak', 'Evening net-load peak (19:29 CDT)')),
    snapshot(LIVE + 'res-ancillary-service-capacity-monitor.json', REL_LIVE + 'res-ancillary-service-capacity-monitor.json',
             '2026-09-26T04:06:47Z', ('late', 'Late evening (23:06 CDT)')),
]

# Time series: dashboard ascapmon (~8-10 s) -> last sample per clock minute. Two windows: cached 17:23-19:23, live 21:06-23:06.
series = []
for path, rel, ret in [(CACHE + 'ercot/db_ancillary-services.json', REL_CACHE + 'ercot/db_ancillary-services.json', '2026-09-26T00:25:12Z'),
                       (LIVE + 'res-ancillary-services.json', REL_LIVE + 'res-ancillary-services.json', '2026-09-26T04:06:50Z')]:
    raw = json.load(open(path)); by_min = {}
    for x in raw['ascapmon']:
        by_min[x['timestamp'][11:16]] = x
    rows = [[k, x['rrs'], x['ecrs'], x['nsrs'], x['deployedRegUp'] + x['undeployedRegUp'], x['deployedRegUp'],
             x['deployedRegDown'] + x['undeployedRegDown'], x['deployedRegDown']] for k, x in sorted(by_min.items())]
    series.append({'source': rel, 'retrievedUtc': ret, 'lastUpdated': raw['lastUpdated'], 'rawPoints': len(raw['ascapmon']),
                   'downsample': 'last ascapmon sample in each clock minute', 'rows': rows})

# ADER Monthly Report (every month in the workbook, read directly; no hand-copied numbers)
import openpyxl
_wb = openpyxl.load_workbook(CACHE + 'bp-data-ingest/ader_monthly.xlsx', read_only=True, data_only=True)
ader_rows = [list(r) for r in _wb['Summary of Commercial ADERs'].iter_rows(min_row=2, values_only=True) if r and r[0]]
assert ader_rows[0][0] == '2025-05' and ader_rows[-1][0] == '2026-08', ader_rows

# ADER Limits of Participation Tracking (as of 06-01-26): approved MW, total and by settlement load zone
_lt = openpyxl.load_workbook(LIVE + 'res-ader-limits-tracking-20260601.xlsx', read_only=True, data_only=True)['Limits & Participation Tracking']
_rows = [r for r in _lt.iter_rows(values_only=True) if any(v is not None for v in r)]
_tot = next(r for r in _rows if 'Total' in r)
_tot = [v for v in _tot if v is not None][1:]
_zones = [v for v in next(r for r in _rows if 'LZ_AEN' in r) if v is not None]
assert len(_tot) == 3 + 3 * len(_zones), (_tot, _zones)
ader_limits = {
    'status': 'REAL', 'asOf': '2026-06-01',
    'source': REL_LIVE + 'res-ader-limits-tracking-20260601.xlsx (https://www.ercot.com/files/docs/2025/05/06/Limits-of-Participation-Tracking_06-01-2026.xlsx, retrieved 2026-09-26T05:29Z)',
    'capsMW': {'capacity': 500, 'nonSpin': 100, 'ecrs': 100}, 'perQseShareMax': 0.9,
    'approvedMW': {'energy': round(_tot[0], 1), 'nonSpin': round(_tot[1], 1), 'ecrs': round(_tot[2], 1)},
    'byZone': {z: {'energy': round(_tot[3 + 3 * i], 1), 'nonSpin': round(_tot[4 + 3 * i], 1), 'ecrs': round(_tot[5 + 3 * i], 1)}
               for i, z in enumerate(_zones)}}

# ESR fleet energy output around each snapshot (energy-storage-resources.json, 5-min; read-only reuse of item load's raw fetch)
_esr = json.load(open(LIVE + 'load-energy-storage-resources.json'))['currentDay']['data']
for s in snaps:
    t = s['asOf'][:16].replace('T', ' ')
    pt = [x for x in _esr if x['timestamp'][:16] <= t][-1]
    s['esrEnergy'] = {'status': 'REAL', 'timestamp': pt['timestamp'], 'dischargingMW': round(pt['totalDischarging']),
                      'chargingMW': round(pt['totalCharging']), 'netOutputMW': round(pt['netOutput']),
                      'source': REL_LIVE + 'load-energy-storage-resources.json (https://www.ercot.com/api/1/services/read/dashboards/energy-storage-resources.json, retrieved 2026-09-26T04:19:09Z)'}

ercot = {
    'status': 'REAL',
    'date': '2026-09-25',
    'footnotes': FOOTNOTES,
    'snapshots': snaps,
    'series': {'fields': ['clockCDT', 'rrsRtAwardMW', 'ecrsRtAwardMW', 'nsrsSparklineMW', 'regUpDeployedPlusUndeployedMW', 'regUpDeployedMW',
                          'regDownDeployedPlusUndeployedMW', 'regDownDeployedMW'],
               'note': 'ERCOT ancillary-services.json keeps only the last 2 h, so the 19:23-21:06 gap is not recoverable from this feed.',
               'windows': series},
    'hourly': {'hourEnding': list(range(1, 25)),
               'planMW': {k: [plan[k][h] for h in range(1, 25)] for k in ['REGUP', 'REGDN', 'RRS', 'ECRS', 'NSPIN']},
               'damMcpcUsdPerMWh': {k: [mcpc[k][h] for h in range(1, 25)] for k in ['REGUP', 'REGDN', 'RRS', 'ECRS', 'NSPIN']},
               'sources': {'plan': REL_LIVE + 'res-np4-33-asplan-pub20260925/ (NP4-33-CD, MIS RTID 12316, doc 1278738466, published 2026-09-25 05:00 CDT)',
                           'mcpc': REL_LIVE + 'res-np4-188-dam-mcpc-od20260925/ (NP4-188-CD, MIS RTID 12329, doc 1278479647, published 2026-09-24 12:51 CDT)'}},
    'ader': {'status': 'REAL', 'source': REL_CACHE + 'bp-data-ingest/ader_monthly.xlsx (ERCOT ADER Monthly Report, https://www.ercot.com/files/docs/2025/04/28/ADER_Monthly_Report_202506_202606.xlsx, sheet "Summary of Commercial ADERs"; cached 2026-09-25 20:14 CDT)',
             'fields': ['month', 'commercialADERs', 'qualifiedEnergyMW', 'qualifiedNonSpinMW', 'qualifiedEcrsMW'], 'rows': ader_rows,
             'pilotCaps': {'totalMW': 500, 'nonSpinMW': 100, 'ecrsMW': 100, 'perQseShare': 0.9,
                           'source': REL_LIVE + 'res-src-ader-gov-doc-3.3.docx (ADER Governing Document Phase 3.3)'},
             'limits': ader_limits},
    'imm2025EsrShare': {'status': 'REAL (IMM annual average, not 25 Sep)', 'REGUP': 0.94, 'REGDN': 0.86, 'RRS': 0.51, 'ECRS': 0.42, 'NSPIN': 0.24,
                        'source': 'Potomac Economics, 2025 State of the Market Report for ERCOT (cached text evidence/scratchpad-20260925/som2025.txt, lines 2912-2914)'},
}
json.dump(ercot, open(sys.argv[1], 'w'))
for s in snaps:
    print(s['label'], 'HE', s['planHourEnding'])
    for p in s['products']:
        print('  ', p['id'], 'plan', p['planMW'], 'award', p['awardMW'], 'cap', p['capabilityMW'], 'gap', p['awardMinusCapabilityMW'],
              'esr', p['esrAwardMW'], p['esrCapabilityMW'], p['esrShareOfAward'], 'mcpc', p['damMcpcUsdPerMWh'])
    print('  ', s['system'])
    print('  ', s['aordc'])
for w in series:
    print(w['source'], len(w['rows']), w['rows'][0], w['rows'][-1])
